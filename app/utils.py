"""Bounded, in-memory resume extraction; original files are never persisted."""

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from flask import current_app
from pypdf import PdfReader
from werkzeug.utils import secure_filename


class ExtractionError(ValueError):
    """A user-correctable upload or text extraction problem."""


def extract_resume(file):
    """Validate extension, size, container, and content before returning text."""
    filename = secure_filename(file.filename or "")
    extension = Path(filename).suffix.lower().lstrip(".")
    if not filename or extension not in current_app.config["ALLOWED_EXTENSIONS"]:
        raise ExtractionError("Choose a PDF, DOCX, or TXT resume.")
    data = file.stream.read(current_app.config["MAX_FILE_SIZE"] + 1)
    if not data:
        raise ExtractionError("Your resume file is empty.")
    if len(data) > current_app.config["MAX_FILE_SIZE"]:
        raise ExtractionError("Your resume must be smaller than 5 MB.")
    try:
        if extension == "txt":
            if b"\x00" in data:
                raise ExtractionError(
                    "This TXT file contains binary data. Export it as UTF-8 text."
                )
            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError as error:
                raise ExtractionError(
                    "Save your TXT resume using UTF-8 encoding."
                ) from error
        elif extension == "pdf":
            if not data.startswith(b"%PDF-"):
                raise ExtractionError("This file is not a valid PDF.")
            reader = PdfReader(BytesIO(data))
            if reader.is_encrypted:
                raise ExtractionError("Remove the PDF password before uploading.")
            if len(reader.pages) > 50:
                raise ExtractionError("Choose a resume with 50 pages or fewer.")
            pages = []
            total = 0
            for page in reader.pages:
                contents = page.get_contents()
                if contents and len(contents.get_data()) > 10 * 1024 * 1024:
                    raise ExtractionError(
                        "This PDF page is too complex. Export a simpler PDF."
                    )
                page_text = page.extract_text() or ""
                total += len(page_text)
                if total > current_app.config["MAX_TEXT_LENGTH"]:
                    raise ExtractionError(
                        "Resume text is too long. Limit it to 60,000 characters."
                    )
                pages.append(page_text)
            text = "\n".join(pages)
        else:
            with ZipFile(BytesIO(data)) as archive:
                entries = archive.infolist()
                if "word/document.xml" not in archive.namelist():
                    raise ExtractionError("This file is not a valid DOCX document.")
                if (
                    len(entries) > 500
                    or sum(item.file_size for item in entries) > 20 * 1024 * 1024
                ):
                    raise ExtractionError(
                        "This DOCX expands beyond the document size limit."
                    )
            document = Document(BytesIO(data))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            cells = [
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ]
            text = "\n".join(paragraphs + cells)
    except ExtractionError:
        raise
    except Exception as error:
        current_app.logger.info("Resume extraction rejected: %s", type(error).__name__)
        raise ExtractionError(
            "We could not read that resume. Re-export it as PDF, DOCX, or UTF-8 TXT."
        ) from error
    text = text.replace("\x00", "").strip()
    if not text:
        raise ExtractionError(
            "No readable text found. Scanned PDFs need OCR; upload a text-based resume."
        )
    if len(text) > current_app.config["MAX_TEXT_LENGTH"]:
        raise ExtractionError("Resume text is too long. Limit it to 60,000 characters.")
    return filename[:255], text
