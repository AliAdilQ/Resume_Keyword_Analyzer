"""Every test uses its own temporary SQLite database, never the demo DB."""

from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app import create_app
from app.extensions import db
from app.models import User

RESUME = "Skills: Python Flask SQL Git. Built reliable backend services and REST API. Improved performance by 30%."
JOB = "Python developer building Flask backend services with SQL, Git, Docker and REST API. Write reliable code and automated testing."


@pytest.fixture
def app(tmp_path):
    application = create_app(
        {
            "TESTING": True,
            "APP_ENV": "development",
            "SECRET_KEY": "test-only-secret",
            "SQLALCHEMY_DATABASE_URI": "sqlite:///" + str(tmp_path / "test.db"),
            "UPLOAD_FOLDER": str(tmp_path / "uploads"),
            "WTF_CSRF_ENABLED": False,
            "LOGIN_MAX_ATTEMPTS": 50,
        }
    )
    with application.app_context():
        for username, admin in (("tester", False), ("other", False), ("admin", True)):
            user = User(
                username=username, email=f"{username}@example.com", is_admin=admin
            )
            user.set_password("Testpass123")
            db.session.add(user)
        db.session.commit()
    yield application
    with application.app_context():
        db.session.remove()
        db.drop_all()
        db.engine.dispose()


@pytest.fixture
def client(app):
    return app.test_client()


@pytest.fixture
def login(client):
    def sign_in(username="tester"):
        return client.post(
            "/login", data={"identity": username, "password": "Testpass123"}
        )

    return sign_in


def pdf_bytes(text=RESUME):
    """Generate an actual text-based PDF without another runtime dependency."""
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): writer._add_object(font)}
            )
        }
    )
    stream = DecodedStreamObject()
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream.set_data(f"BT /F1 12 Tf 30 700 Td ({escaped}) Tj ET".encode("ascii"))
    page[NameObject("/Contents")] = writer._add_object(stream)
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def docx_bytes(text=RESUME):
    document = Document()
    document.add_paragraph(text)
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "GitHub Linux"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
